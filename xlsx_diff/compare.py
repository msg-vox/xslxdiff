"""Read-only workbook comparison; formulas are never evaluated."""
import json
from dataclasses import dataclass, asdict
from datetime import date, datetime, time, timedelta
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.xml.functions import tostring


@dataclass(frozen=True)
class Change:
    sheet: str
    reference: str
    change_type: str
    property: str
    old_value: str
    new_value: str

    def as_dict(self):
        return asdict(self)


def display(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)


def value_of(cell):
    if cell is None or isinstance(cell, MergedCell):
        return None, "n"
    value = cell.value
    # openpyxl represents array/data-table formulas as objects, not strings.
    if cell.data_type == "f" and not isinstance(value, str):
        value = dict(vars(value))
    return value, cell.data_type


def style_of(cell, name):
    if name == "number_format":
        return cell.number_format
    return tostring(getattr(cell, name).to_tree()).decode("utf-8")


def compare_workbooks(old_path, new_path, *, ignore_formatting=False,
                      ignore_blank=False, values_only=False):
    """Return changes. Blank means None or an empty string on both sides.

    Sparse cell storage avoids scanning enormous unused worksheet rectangles.
    It includes explicitly styled blank cells. Inputs are never saved.
    """
    old = load_workbook(old_path, data_only=False)
    try:
        new = load_workbook(new_path, data_only=False)
        try:
            return _compare(old, new, ignore_formatting, ignore_blank, values_only)
        finally:
            new.close()
    finally:
        old.close()


def _compare(old, new, ignore_formatting, ignore_blank, values_only):
    changes = []

    def add(sheet, ref, prop, before, after, kind=None):
        if before == after and type(before) is type(after):
            return
        if kind is None:
            kind = "added" if before is None else "removed" if after is None else "modified"
        changes.append(Change(sheet, ref, kind, prop, display(before), display(after)))

    old_names, new_names = old.sheetnames, new.sheetnames
    common_old = [s for s in old_names if s in new_names]
    common_new = [s for s in new_names if s in old_names]
    blank_book = Workbook()
    blank_cell = blank_book.active.cell(1, 1)
    for name in old_names + [s for s in new_names if s not in old_names]:
        a = old[name] if name in old_names else None
        b = new[name] if name in new_names else None
        if a is None:
            add(name, "", "worksheet", None, "present")
        elif b is None:
            add(name, "", "worksheet", "present", None)
        elif common_old != common_new and common_old.index(name) != common_new.index(name):
            add(name, "", "worksheet_order", old_names.index(name) + 1,
                new_names.index(name) + 1, "modified")
        acells = a._cells if a is not None else {}
        bcells = b._cells if b is not None else {}
        for row, col in sorted(acells.keys() | bcells.keys()):
            ca, cb = acells.get((row, col)), bcells.get((row, col))
            av, at = value_of(ca)
            bv, bt = value_of(cb)
            if ignore_blank and (av is None or av == "") and (bv is None or bv == ""):
                continue
            ref = f"{get_column_letter(col)}{row}"
            prop = "formula" if at == "f" or bt == "f" else "value"
            if av != bv or type(av) is not type(bv):
                add(name, ref, prop, av, bv)
            elif at != bt and av is not None:
                add(name, ref, "cell_type", at, bt, "modified")
            if not (ignore_formatting or values_only):
                for part in ("font", "fill", "border", "alignment", "number_format"):
                    add(name, ref, part, style_of(ca if ca is not None else blank_cell, part),
                        style_of(cb if cb is not None else blank_cell, part), "modified")
        if values_only:
            continue
        amerge = {str(r) for r in a.merged_cells.ranges} if a is not None else set()
        bmerge = {str(r) for r in b.merged_cells.ranges} if b is not None else set()
        for ref in sorted(amerge - bmerge):
            add(name, ref, "merged_range", ref, None)
        for ref in sorted(bmerge - amerge):
            add(name, ref, "merged_range", None, ref)
        _dimensions(a, b, name, add)
    return changes


def _dimensions(a, b, name, add):
    def rows(sheet):
        if sheet is None:
            return {}
        return {i: (bool(d.hidden), d.height) for i, d in sheet.row_dimensions.items()}

    def columns(sheet):
        result = {}
        if sheet is not None:
            for key, d in sheet.column_dimensions.items():
                from openpyxl.utils import column_index_from_string
                start = d.min or column_index_from_string(key)
                end = d.max or start
                for i in range(start, end + 1):
                    result[i] = (bool(d.hidden), d.width)
        return result

    def defaults(sheet):
        if sheet is None:
            from openpyxl.worksheet.dimensions import SheetFormatProperties
            return SheetFormatProperties()
        return sheet.sheet_format

    da, db = defaults(a), defaults(b)
    for prop in ("defaultRowHeight", "defaultColWidth", "baseColWidth", "zeroHeight"):
        add(name, "", prop, getattr(da, prop), getattr(db, prop), "modified")
    for getter, props, reference, default_a, default_b in (
        (rows, ("row_hidden", "row_height"), lambda i: f"{i}:{i}",
         (bool(da.zeroHeight), da.defaultRowHeight), (bool(db.zeroHeight), db.defaultRowHeight)),
        (columns, ("column_hidden", "column_width"), lambda i: f"{get_column_letter(i)}:{get_column_letter(i)}",
         (False, da.defaultColWidth or 13.0), (False, db.defaultColWidth or 13.0)),
    ):
        before, after = getter(a), getter(b)
        for i in sorted(before.keys() | after.keys()):
            va, vb = before.get(i, default_a), after.get(i, default_b)
            # Unspecified row heights inherit the sheet default.
            va = (va[0], default_a[1] if va[1] is None else va[1])
            vb = (vb[0], default_b[1] if vb[1] is None else vb[1])
            for prop, x, y in zip(props, va, vb):
                add(name, reference(i), prop, x, y, "modified")
