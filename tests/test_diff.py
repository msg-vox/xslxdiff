import contextlib
import csv
import hashlib
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.formula import ArrayFormula

from xlsx_diff.cli import main
from xlsx_diff.compare import compare_workbooks


class DiffTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.old_path, self.new_path = self.root / 'old.xlsx', self.root / 'new.xlsx'
        self.old, self.new = Workbook(), Workbook()

    def diff(self, **kwargs):
        self.old.save(self.old_path)
        self.new.save(self.new_path)
        return compare_workbooks(self.old_path, self.new_path, **kwargs)

    def test_empty_and_identical(self):
        self.assertEqual(self.diff(), [])
        self.old.active['A1'] = self.new.active['A1'] = '=SUM(B1:B4)'
        self.assertEqual(self.diff(), [])

    def test_values_formulas_types_and_dimensions(self):
        a, b = self.old.active, self.new.active
        a['A1'], b['A1'] = 'old', 'new'
        a['B2'] = 'gone'
        b['Z100'] = 'added'
        a['C3'], b['C3'] = '=SUM(A1:A2)', '=SUM(A1:A3)'
        a['D4'], b['D4'] = 1, True
        changes = {c.reference: c for c in self.diff()}
        self.assertEqual(changes['A1'].change_type, 'modified')
        self.assertEqual(changes['B2'].change_type, 'removed')
        self.assertEqual(changes['Z100'].change_type, 'added')
        self.assertEqual(changes['C3'].property, 'formula')
        self.assertEqual(changes['C3'].old_value, '=SUM(A1:A2)')
        self.assertEqual(changes['C3'].new_value, '=SUM(A1:A3)')
        self.assertEqual(changes['D4'].new_value, 'true')

    def test_formula_objects_and_literal_formula_text(self):
        self.old.active['A1'] = ArrayFormula(ref='A1:A3', text='=B1:B3*2')
        self.new.active['A1'] = ArrayFormula(ref='A1:A3', text='=B1:B3*3')
        self.old.active['C1'] = '=1+1'
        self.new.active['C1'] = '=1+1'
        self.new.active['C1'].data_type = 's'
        changes = self.diff()
        self.assertTrue(any(c.property == 'formula' and '*3' in c.new_value for c in changes))
        self.assertTrue(any(c.property == 'cell_type' for c in changes))

    def test_all_style_components_and_blank_filter(self):
        cell = self.new.active['A1']
        cell.font = Font(bold=True, color='FF0000')
        cell.fill = PatternFill('solid', fgColor='00FF00')
        cell.border = Border(left=Side(style='thin'))
        cell.alignment = Alignment(horizontal='center')
        cell.number_format = '0.00'
        expected = {'font', 'fill', 'border', 'alignment', 'number_format'}
        self.assertEqual({c.property for c in self.diff()}, expected)
        self.assertEqual(self.diff(ignore_blank=True), [])
        self.assertEqual(self.diff(ignore_formatting=True), [])
        self.assertEqual(self.diff(values_only=True), [])
        self.new.active['B1'] = 'new'
        self.assertEqual([c.reference for c in self.diff(ignore_blank=True)], ['B1'])

    def test_merges_hidden_rows_columns_and_sizes(self):
        sheet = self.new.active
        sheet.merge_cells('C3:D4')
        sheet.row_dimensions[5].hidden = True
        sheet.row_dimensions[6].height = 30
        sheet.column_dimensions.group('E', 'G', hidden=True)
        sheet.column_dimensions['J'].width = 25
        changes = self.diff(ignore_formatting=True)
        props = {c.property for c in changes}
        self.assertEqual(props, {'merged_range', 'row_hidden', 'row_height', 'column_hidden', 'column_width'})
        self.assertEqual({c.reference for c in changes if c.property == 'column_hidden'}, {'E:E', 'F:F', 'G:G'})
        self.assertEqual(self.diff(values_only=True), [])

    def test_merge_and_dimensions_removed(self):
        self.old.active.merge_cells('A2:B2')
        self.old.active.row_dimensions[2].height = 22
        self.old.active.column_dimensions['C'].hidden = True
        changes = self.diff()
        self.assertTrue(any(c.property == 'merged_range' and c.change_type == 'removed' for c in changes))
        self.assertTrue(any(c.property == 'row_height' and c.new_value == '15.0' for c in changes))
        self.assertTrue(any(c.property == 'column_hidden' and c.new_value == 'false' for c in changes))

    def test_sheet_defaults(self):
        self.new.active.sheet_format.defaultRowHeight = 22
        self.new.active.sheet_format.defaultColWidth = 18
        changes = self.diff()
        self.assertEqual({c.property for c in changes}, {'defaultRowHeight', 'defaultColWidth'})

    def test_added_removed_reordered_sheets(self):
        for book in (self.old, self.new):
            book.create_sheet('Second')
        self.old.create_sheet('Deleted')['A1'] = 'gone'
        self.new.create_sheet('Added')['B1'] = 'new'
        self.new.move_sheet('Second', offset=-1)
        changes = self.diff()
        self.assertEqual(len([c for c in changes if c.property == 'worksheet_order']), 2)
        self.assertTrue(any(c.sheet == 'Deleted' and c.reference == 'A1' for c in changes))
        self.assertTrue(any(c.sheet == 'Added' and c.reference == 'B1' for c in changes))
        self.assertEqual(len([c for c in changes if c.property == 'worksheet']), 2)

    def test_inserted_sheet_does_not_reorder_common_sheets(self):
        for book in (self.old, self.new):
            book.create_sheet('Second')
        self.new.create_sheet('First', 0)
        self.assertFalse(any(c.property == 'worksheet_order' for c in self.diff()))

    def test_reports_and_original_hashes(self):
        self.new.active.title = '<Sheet&>'
        self.old.active.title = '<Sheet&>'
        self.old.active['A1'] = 'before,\ncomma'
        self.new.active['A1'] = '<script>alert("x")</script>'
        self.diff()
        before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (self.old_path, self.new_path)]
        report = self.root / 'result.html'
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            result = main([str(self.old_path), str(self.new_path), '--output', str(report)])
        self.assertEqual(result, 0)
        self.assertIn('1 changes', out.getvalue())
        html = report.read_text()
        self.assertIn('&lt;script&gt;', html)
        self.assertNotIn('<script>', html)
        self.assertIn('&lt;Sheet&amp;&gt;', html)
        self.assertIn('class="modified"', html)
        with report.with_suffix('.csv').open(newline='') as stream:
            records = list(csv.DictReader(stream))
        self.assertEqual(records[0]['old_value'], 'before,\ncomma')
        self.assertEqual(records[0]['new_value'], '<script>alert("x")</script>')
        self.assertEqual(before, [hashlib.sha256(p.read_bytes()).hexdigest() for p in (self.old_path, self.new_path)])

    def test_errors_and_input_output_aliases(self):
        self.diff()
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(main(['missing.xlsx', str(self.new_path)]), 2)
            self.assertEqual(main([str(self.old_path), str(self.new_path), '--output', str(self.old_path)]), 2)
            alias = self.root / 'alias.html'
            alias.symlink_to(self.old_path)
            self.assertEqual(main([str(self.old_path), str(self.new_path), '--output', str(alias)]), 2)
            alias.unlink()
            alias.hardlink_to(self.old_path)
            self.assertEqual(main([str(self.old_path), str(self.new_path), '--output', str(alias)]), 2)
            self.assertEqual(main([str(self.old_path), str(self.new_path), '--output', str(self.root / 'same.csv')]), 2)
            self.new_path.write_text('not a zip')
            self.assertEqual(main([str(self.old_path), str(self.new_path)]), 2)

    def test_sparse_far_corner(self):
        self.new.active['XFD1048576'] = 'far'
        changes = self.diff(values_only=True)
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0].reference, 'XFD1048576')


if __name__ == '__main__':
    unittest.main()
